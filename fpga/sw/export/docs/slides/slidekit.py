"""slidekit.py — build lecture decks in the UEC / Pham Lab template.

template_uec.pptx holds six slides cut from the course's reference deck
("04 Creating RISC-V Project with Chipyard"); every slide of a new deck is a
copy of one of them, so logos, header bar, watermark and page number keep the
reference look. Content is added as native shapes (text, code boxes, tables,
charts, diagrams), so every slide stays editable in PowerPoint.

    from slidekit import Deck
    d = Deck()
    d.title("「 Lecture 5 」 ...", author="[Author]", date="[YYYY/MM]")
    s = d.content("2. SoC Configuration", "(1/7) Config chain")
    d.text(s, ["...", (1, "sub item")], x=0.4, y=0.9, w=12.5, h=3)
    d.save("05_....pptx")
"""
import copy
import os
import re

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "template_uec.pptx")
T_TITLE, T_PURPOSE, T_OUTLINE, T_CONTENT, T_HIGHLIGHT, T_THANKS = range(6)

FONT = "Times New Roman"
MONO = "Consolas"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(0x80, 0x80, 0x80)
GREEN = RGBColor(0x37, 0x86, 0x3C)      # outline highlight
RED = RGBColor(0xC0, 0x00, 0x00)
CODE_BG = RGBColor(0xF4, 0xF4, 0xF4)
CODE_FG = RGBColor(0x1F, 0x1F, 0x1F)
HEAD_BG = RGBColor(0xC5, 0xE0, 0xB4)    # same tint as the header bar
BOX_BG = RGBColor(0xFF, 0xF2, 0xCC)


def _style(run, size, bold=None, color=BLACK, font=FONT, italic=None):
    f = run.font
    f.name = font
    f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def _rich(par, text, size, color=BLACK, bold=False):
    """**bold**, `mono` and *italic* inside one paragraph."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not tok:
            continue
        r = par.add_run()
        if tok.startswith("**"):
            r.text = tok[2:-2]
            _style(r, size, True, color)
        elif tok.startswith("`"):
            r.text = tok[1:-1]
            _style(r, size - 2, bold, RGBColor(0x8B, 0x1A, 0x1A), MONO)
        elif tok.startswith("*") and len(tok) > 2:
            r.text = tok[1:-1]
            _style(r, size, bold, color, italic=True)
        else:
            r.text = tok
            _style(r, size, bold, color)


def _bullet(par, level):
    """Real bullets (• then –) with a hanging indent."""
    pPr = par._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.28 + 0.32 * level))))
    pPr.set("indent", str(int(-Inches(0.24))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if level == 0 else "–"}))


def _clear(tf):
    """Leave one empty paragraph (keeps its paragraph properties)."""
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    return tf.paragraphs[0]


class Deck:
    def __init__(self, template=TEMPLATE):
        self.prs = Presentation(template)
        self._tpl = list(self.prs.slides)
        self.sections = []

    # ---- slide cloning -------------------------------------------------
    def _clone(self, kind):
        src = self._tpl[kind]
        dst = self.prs.slides.add_slide(src.slide_layout)
        for sp in list(dst.shapes):
            sp._element.getparent().remove(sp._element)
        rmap = {}
        for rid, rel in src.part.rels.items():
            if "image" in rel.reltype:
                rmap[rid] = dst.part.relate_to(rel._target, rel.reltype)
        for el in src.shapes._spTree.iterchildren():
            if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
                continue
            new = copy.deepcopy(el)
            for node in new.iter():
                for attr in (qn("r:embed"), qn("r:link")):
                    if node.get(attr) in rmap:
                        node.set(attr, rmap[node.get(attr)])
            dst.shapes._spTree.append(new)
        n = len(self.prs.slides) - len(self._tpl)
        for sh in dst.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip().isdigit():
                self._set_text(sh, str(n))
        return dst

    @staticmethod
    def _set_text(shape, text):
        runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
        if not runs:            # slide-number field: set its cached text too,
            for t in shape._element.iter(qn("a:t")):   # some viewers show it as is
                t.text = text
            return
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""

    @staticmethod
    def _find(slide, contains):
        for sh in slide.shapes:
            if sh.has_text_frame and contains in sh.text_frame.text:
                return sh
        raise KeyError(contains)

    # ---- slide kinds ---------------------------------------------------
    def title(self, lecture, course="「Course」\n[Course name]", author="[Author]",
              date="[YYYY/MM]"):
        s = self._clone(T_TITLE)
        c = self._find(s, "Course")
        for p, line in zip(c.text_frame.paragraphs, course.split("\n")):
            p.runs[0].text = line
            for r in p.runs[1:]:
                r.text = ""
        lec = self._find(s, "Lecture")
        m = re.match(r"(「[^」]*」)\s*(.*)", lecture)
        runs = lec.text_frame.paragraphs[0].runs
        runs[0].text = m.group(1) if m else lecture
        for r in runs[1:]:
            r.text = ""
        if m and len(runs) > 1:
            runs[-1].text = "  " + m.group(2)
        self._set_text(self._find(s, "Hoang"), author)
        self._set_text(self._find(s, "2022"), date)
        return s

    def purpose(self, heading, items, footer=None):
        s = self._clone(T_PURPOSE)
        tf = self._find(s, "purpose").text_frame
        p0 = _clear(tf)
        _rich(p0, heading, 24, bold=True)
        for it in items:
            par = tf.add_paragraph()
            _bullet(par, 0)
            _rich(par, it, 22)
        if footer:
            par = tf.add_paragraph()
            _rich(par, footer, 18, GRAY)
        return s

    def outline(self, current=None):
        s = self._clone(T_OUTLINE)
        tf = self._find(s, "Chipyard organization").text_frame
        tmpl = copy.deepcopy(tf.paragraphs[0]._p)
        for p in list(tf.paragraphs):
            p._p.getparent().remove(p._p)
        for name in self.sections:
            tf._txBody.append(copy.deepcopy(tmpl))
        size = 32 if len(self.sections) <= 5 else 28
        for i, (par, name) in enumerate(zip(tf.paragraphs, self.sections)):
            for r in par.runs[1:]:
                r._r.getparent().remove(r._r)
            r = par.runs[0]
            r.text = name
            if current is None:
                _style(r, size, False, BLACK)
            else:
                _style(r, size, i == current, GREEN if i == current else GRAY)
        return s

    def content(self, section, sub):
        s = self._clone(T_CONTENT)
        for sh in list(s.shapes):
            keep = sh.has_text_frame and (sh.text_frame.text.strip().isdigit()
                                          or "SoC Generation" in sh.text_frame.text)
            if not keep:
                sh._element.getparent().remove(sh._element)
        runs = self._find(s, "SoC Generation").text_frame.paragraphs[0].runs
        runs[0].text = section + " "
        runs[1].text = sub
        n = len(section) + 1 + len(sub)
        size = max(24, min(36, int(13.0 * 72 / (n * 0.5))))   # one line, 13" bar
        for r in runs[:2]:
            r.font.size = Pt(size)
        return s

    def highlight(self, title, items, next_title=None, next_text=None):
        s = self._clone(T_HIGHLIGHT)
        tf = self._find(s, "highlight").text_frame
        p0 = _clear(tf)
        p0.alignment = PP_ALIGN.CENTER
        _rich(p0, f"*{title}*", 22, bold=True)
        for it in items:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            _bullet(p, 0)
            _rich(p, it, 19)
        if next_title:
            for txt, b in ((f"*{next_title}*", True), (next_text, False)):
                p = tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                _rich(p, txt, 20, bold=b)
        return s

    def thanks(self, date="[YYYY/MM]"):
        s = self._clone(T_THANKS)
        self._set_text(self._find(s, "2022"), date)
        return s

    # ---- content helpers -----------------------------------------------
    def text(self, slide, items, x, y, w, h, size=20, bullets=True, color=BLACK,
             anchor=MSO_ANCHOR.TOP):
        """items: str (level 0), (level, str), or ('', str) for a plain line."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, it in enumerate(items):
            level, txt = it if isinstance(it, tuple) else (0, it)
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.space_after = Pt(4)
            if bullets and level != "":
                _bullet(par, level)
            _rich(par, txt, size - (2 if level == 1 else 0), color)
        return tb

    def code(self, slide, code, x, y, w, h, size=14, title=None):
        """Monospace box. The font shrinks so the longest line and all lines fit
        w x h; the box then shrinks to its content."""
        lines = code.strip("\n").split("\n") + ([title] if title else [])
        cols = max(len(l) for l in lines)
        size = min(size, (w - 0.2) * 72 / (cols * 0.61), (h - 0.2) * 72 / (len(lines) * 1.22))
        size = max(size, 8)
        h = min(h, len(lines) * size * 1.22 / 72 + 0.22)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = CODE_BG
        box.line.color.rgb = RGBColor(0xC8, 0xC8, 0xC8)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(0.08))
        lines = code.strip("\n").split("\n")
        if title:
            lines = [None] + lines
        for i, line in enumerate(lines):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = PP_ALIGN.LEFT
            r = par.add_run()
            if line is None:
                r.text = title
                _style(r, size, True, GRAY, MONO)
                continue
            r.text = line
            par.line_spacing = 1.0
            comment = line.lstrip().startswith(("//", "#", "/*", "*"))
            _style(r, size, None, RGBColor(0x2E, 0x7D, 0x32) if comment else CODE_FG, MONO)
        return box

    def console(self, slide, text, x, y, w, h, size=12, title=None):
        """Real board console output: dark terminal box."""
        box = self.code(slide, text, x, y, w, h, size, title)
        box.fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
        for i, par in enumerate(box.text_frame.paragraphs):
            for r in par.runs:
                t = r.text
                r.font.color.rgb = (RGBColor(0xAA, 0xAA, 0xAA) if title and i == 0 else
                                    RGBColor(0x8B, 0xE9, 0x8B) if "PASS" in t else
                                    RGBColor(0xFF, 0x80, 0x80) if "FAIL" in t else
                                    RGBColor(0xE6, 0xE6, 0xE6))
        return box

    def table(self, slide, rows, x, y, w, h, size=16, col_w=None, head=True):
        gt = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                    Inches(w), Inches(h))
        t = gt.table
        if col_w:
            for i, cw in enumerate(col_w):
                t.columns[i].width = Inches(cw)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.06)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.text_frame.word_wrap = True
                _rich(cell.text_frame.paragraphs[0], str(val), size, bold=(head and r == 0))
                cell.fill.solid()
                cell.fill.fore_color.rgb = HEAD_BG if head and r == 0 else (
                    RGBColor(0xFF, 0xFF, 0xFF) if r % 2 else RGBColor(0xF2, 0xF8, 0xEE))
        return gt

    def photo(self, slide, label, x, y, w, h):
        """Placeholder for a photo still to be taken."""
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                     Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xE7, 0xE6, 0xE6)
        box.line.color.rgb = GRAY
        box.line.dash_style = 4
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(["PHOTO"] + label.split("\n")):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = PP_ALIGN.CENTER
            _rich(par, line, 16, GRAY, i == 0)
        return box

    def box(self, slide, label, x, y, w, h, fill=BOX_BG, size=16, bold=False,
            line=GRAY, color=BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        b = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        b.fill.solid()
        b.fill.fore_color.rgb = fill
        b.line.color.rgb = line
        b.shadow.inherit = False
        tf = b.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(0.04))
        for i, line_txt in enumerate(label.split("\n")):
            par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            par.alignment = PP_ALIGN.CENTER
            _rich(par, line_txt, size, color, bold)
        return b

    def arrow(self, slide, x1, y1, x2, y2, color=GRAY, both=False, width=2):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                       Inches(x2), Inches(y2))
        c.line.color.rgb = color
        c.line.width = Pt(width)
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
        if both:
            ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle"}))
        return c

    def bar_chart(self, slide, categories, series, x, y, w, h, stacked=True,
                  number_format="#,##0", size=12):
        data = CategoryChartData()
        data.categories = categories
        for name, vals in series:
            data.add_series(name, vals)
        kind = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
        ch = gf.chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
        ch.legend.font.name = FONT
        ch.category_axis.tick_labels.font.size = Pt(size)
        ch.category_axis.tick_labels.font.name = FONT
        va = ch.value_axis
        va.tick_labels.font.size = Pt(size)
        va.tick_labels.font.name = FONT
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        va.has_major_gridlines = True
        return gf

    @staticmethod
    def below(shape, gap=0.2):
        """y (inches) just under a placed shape, for stacking."""
        return (shape.top + shape.height) / 914400 + gap

    # ---- output --------------------------------------------------------
    def save(self, path):
        ids = self.prs.slides._sldIdLst
        for sid in list(ids)[: len(self._tpl)]:
            self.prs.part.drop_rel(sid.rId)
            ids.remove(sid)
        self.prs.save(path)
